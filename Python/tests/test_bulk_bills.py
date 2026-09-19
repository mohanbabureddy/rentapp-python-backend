import io
import unittest
from datetime import datetime
from types import SimpleNamespace as N

from app.bulk_bills import (
    BulkError, build_template, import_planned, parse_amount, parse_month, plan_rows, read_rows,
)


def _users(*names, role="TENANT"):
    known = {n: N(username=n, role=role) for n in names}
    return N(find_by_username=lambda u: known.get(u))


def _bills(existing=()):
    return N(find_by_tenant_name_and_month=lambda t, m, ty: N() if (t, m, ty) in existing else None)


def _rent(text):
    return read_rows("bills.csv", text.encode("utf-8"), "rent")


def _elec(text):
    return read_rows("bills.csv", text.encode("utf-8"), "electricity")


class ParseTest(unittest.TestCase):
    def test_amounts(self):
        self.assertEqual(parse_amount("1,200"), 1200.0)
        self.assertEqual(parse_amount("Rs. 850"), 850.0)
        self.assertEqual(parse_amount(500), 500.0)
        self.assertIsNone(parse_amount(""))
        self.assertIsNone(parse_amount(None))
        for bad in ("abc", -5, "-1"):
            with self.assertRaises(ValueError):
                parse_amount(bad)

    def test_months(self):
        self.assertEqual(parse_month("2026-09"), "2026-09")
        self.assertEqual(parse_month(datetime(2026, 9, 1)), "2026-09")
        self.assertEqual(parse_month("Sep-2026"), "2026-09")
        with self.assertRaises(ValueError):
            parse_month("September")


class ReadRowsTest(unittest.TestCase):
    def test_rent_template_headings(self):
        rows = _rent("Username,Rent,Water,Miscellaneous\nRoom1,10000,300,200\n")
        self.assertEqual(rows[0][0], 2)
        rec = rows[0][1]
        self.assertEqual((rec["tenant"], rec["rent"], rec["water"], rec["miscellaneous"]),
                         ("Room1", "10000", "300", "200"))

    def test_electricity_template_headings(self):
        rows = _elec("Username,Electricity Bill Amount\nRoom1,850\nRoom2,600\n")
        self.assertEqual(len(rows), 2)
        self.assertEqual(rows[0][1]["electricity"], "850")

    def test_the_two_uploads_cannot_be_mixed(self):
        with self.assertRaises(BulkError) as ctx:
            _rent("Username,Month,Rent,Electricity Bill Amount\nRoom1,2026-09,100,50\n")
        self.assertIn("Electricity column", str(ctx.exception))
        with self.assertRaises(BulkError) as ctx:
            _elec("Username,Electricity Bill Amount,Rent\nRoom1,50,100\n")
        self.assertIn("Rent, Water or Miscellaneous", str(ctx.exception))

    def test_required_columns(self):
        with self.assertRaises(BulkError):
            _rent("Room,Foo\nRoom1,5\n")
        with self.assertRaises(BulkError):
            _rent("Username,Water\nRoom1,5\n")
        with self.assertRaises(BulkError):
            _elec("Username,Month\nRoom1,2026-09\n")

    def test_wrong_file_type_and_kind(self):
        with self.assertRaises(BulkError):
            read_rows("bills.pdf", b"x", "rent")
        with self.assertRaises(BulkError):
            read_rows("bills.csv", b"Username,Rent\nRoom1,1\n", "everything")

    def test_blank_rows_skipped(self):
        self.assertEqual(len(_rent("Username,Rent\nRoom1,100\n,\n\n")), 1)

    def test_real_xlsx_and_the_generated_templates(self):
        from openpyxl import Workbook, load_workbook
        wb = Workbook()
        ws = wb.active
        ws.append(["Username", "Electricity Bill Amount"])
        ws.append(["Room1", 850])
        out = io.BytesIO()
        wb.save(out)
        rows = read_rows("t.xlsx", out.getvalue(), "electricity")
        self.assertEqual((rows[0][1]["tenant"], rows[0][1]["electricity"]), ("Room1", 850))
        headings = {k: [c.value for c in load_workbook(io.BytesIO(build_template(k))).active[1]] for k in ("rent", "electricity")}
        self.assertEqual(headings["rent"], ["Username", "Rent", "Water", "Miscellaneous"])
        self.assertEqual(headings["electricity"], ["Username", "Electricity Bill Amount"])

    def test_month_is_not_a_column_of_either_template(self):
        from openpyxl import load_workbook
        for kind in ("rent", "electricity"):
            headings = [c.value for c in load_workbook(io.BytesIO(build_template(kind))).active[1]]
            self.assertNotIn("Month", headings)


class PlanRentTest(unittest.TestCase):
    def test_rent_row_makes_one_rent_bill_only(self):
        plan = plan_rows(_rent("Username,Rent,Water,Miscellaneous\nRoom1,10000,300,200\n"),
                         "rent", "2026-09", _users("Room1"), _bills())
        self.assertEqual(len(plan), 1)
        p = plan[0]
        self.assertEqual((p["billType"], p["rent"], p["water"], p["miscellaneous"], p["status"]), ("RENT", 10000.0, 300.0, 200.0, "ok"))
        self.assertNotIn("amount", p)
        b = p["_bill"]
        self.assertEqual((b.rent, b.water, b.miscellaneous, b.electricity), (10000.0, 300.0, 200.0, None))

    def test_water_and_misc_are_optional(self):
        plan = plan_rows(_rent("Username,Month,Rent\nRoom1,2026-09,10000\n"), "rent", None, _users("Room1"), _bills())
        p = plan[0]
        self.assertEqual((p["status"], p["rent"], p["water"], p["miscellaneous"]), ("ok", 10000.0, 0.0, 0.0))
        self.assertEqual((p["_bill"].water, p["_bill"].miscellaneous), (0.0, 0.0))  # saved as 0, not empty

    def test_water_only_with_blank_misc_shows_misc_as_zero(self):
        csv_text = "\n".join(["Username,Rent,Water,Miscellaneous", "Room1,10000,750,"])
        plan = plan_rows(_rent(csv_text), "rent", "2026-09", _users("Room1"), _bills())
        p = plan[0]
        self.assertEqual((p["rent"], p["water"], p["miscellaneous"]), (10000.0, 750.0, 0.0))
        self.assertEqual(p["_bill"].miscellaneous, 0.0)

    def test_missing_rent_is_an_error(self):
        plan = plan_rows(_rent("Username,Month,Rent,Water\nRoom1,2026-09,,300\n"), "rent", None, _users("Room1"), _bills())
        self.assertEqual((plan[0]["status"], plan[0]["message"]), ("error", "Rent is empty"))

    def test_month_from_the_picker_when_no_column_value(self):
        plan = plan_rows(_rent("Username,Rent\nRoom1,100\n"), "rent", "2026-09", _users("Room1"), _bills())
        self.assertEqual((plan[0]["month"], plan[0]["status"]), ("2026-09", "ok"))

    def test_no_month_anywhere_is_a_row_error(self):
        plan = plan_rows(_rent("Username,Rent\nRoom1,100\n"), "rent", None, _users("Room1"), _bills())
        self.assertEqual(plan[0]["status"], "error")
        self.assertIn("month", plan[0]["message"].lower())


class PlanElectricityTest(unittest.TestCase):
    def test_electricity_row_makes_one_electricity_bill_only(self):
        plan = plan_rows(_elec("Username,Electricity Bill Amount\nRoom1,850\n"), "electricity", "2026-09", _users("Room1"), _bills())
        self.assertEqual(len(plan), 1)
        b = plan[0]["_bill"]
        self.assertEqual((plan[0]["billType"], plan[0]["month"], plan[0]["electricity"]), ("ELECTRICITY", "2026-09", 850.0))
        self.assertEqual((b.electricity, b.rent, b.water), (850.0, None, None))

    def test_zero_or_empty_amount_is_an_error(self):
        plan = plan_rows(_elec("Username,Electricity Bill Amount\nRoom1,0\nRoom2,\n"), "electricity", "2026-09", _users("Room1", "Room2"), _bills())
        self.assertEqual([p["status"] for p in plan], ["error", "error"])


class PlanCommonTest(unittest.TestCase):
    def test_bad_rows_are_reported_not_fatal(self):
        rows = _rent("Username,Month,Rent\nGhost,2026-09,100\nRoom1,2026-09,abc\nRoom1,2026-13,100\nAdmin1,2026-09,100\n")
        users = N(find_by_username=lambda u: {"Room1": N(role="TENANT"), "Admin1": N(role="ADMIN")}.get(u))
        plan = plan_rows(rows, "rent", None, users, _bills())
        self.assertEqual([p["status"] for p in plan], ["error"] * 4)
        self.assertIn("not found", plan[0]["message"])
        self.assertIn("valid amount", plan[1]["message"])
        self.assertIn("valid month", plan[2]["message"])
        self.assertIn("not a tenant", plan[3]["message"])

    def test_existing_bill_of_the_other_type_is_not_a_duplicate(self):
        existing = {("Room1", "2026-09", "ELECTRICITY")}
        rent_plan = plan_rows(_rent("Username,Month,Rent\nRoom1,2026-09,100\n"), "rent", None, _users("Room1"), _bills(existing))
        self.assertEqual(rent_plan[0]["status"], "ok")

    def test_existing_and_repeated_bills_are_skipped(self):
        rows = _rent("Username,Month,Rent\nRoom1,2026-09,100\nRoom1,2026-09,100\nRoom2,2026-09,100\n")
        plan = plan_rows(rows, "rent", None, _users("Room1", "Room2"), _bills({("Room2", "2026-09", "RENT")}))
        self.assertEqual([p["status"] for p in plan], ["ok", "duplicate", "duplicate"])

    def test_too_many_bills(self):
        body = "Username,Month,Rent\n" + "".join(f"R{i},2026-09,1\n" for i in range(41))
        with self.assertRaises(BulkError):
            plan_rows(_rent(body), "rent", None, _users(*[f"R{i}" for i in range(41)]), _bills())


class ImportTest(unittest.TestCase):
    def test_each_valid_bill_goes_through_add_bill_once(self):
        saved = []

        def add_bill(bill):
            if bill.tenant_name == "Room2":
                raise ValueError("Rent bill already exists for this tenant and month.")
            saved.append((bill.tenant_name, bill.bill_type))

        rows = _rent("Username,Month,Rent\nRoom1,2026-09,100\nRoom2,2026-09,100\nGhost,2026-09,1\n")
        plan = plan_rows(rows, "rent", None, _users("Room1", "Room2"), _bills())
        import_planned(plan, N(add_bill=add_bill))
        self.assertEqual(saved, [("Room1", "RENT")])
        self.assertEqual([p["status"] for p in plan], ["added", "duplicate", "error"])


if __name__ == "__main__":
    unittest.main()
