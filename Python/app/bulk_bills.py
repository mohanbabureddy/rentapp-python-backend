"""Bulk bill upload from an Excel (.xlsx) or CSV file.

Two separate kinds of upload, never mixed:
  * "rent"        -> RENT bills (Rent, Water, Miscellaneous)
  * "electricity" -> ELECTRICITY bills (Electricity Bill Amount)

Flow: read_rows() -> plan_rows() (validation, nothing saved) -> import_planned()
(saves each valid bill through TenantBillService.add_bill, which sends one
email per bill, exactly like adding a bill by hand).
"""
import csv
import io
import math
import re
from datetime import date, datetime
from typing import Any, Dict, List, Optional, Tuple

from app.models import TenantBill

MAX_BYTES = 1_000_000
MAX_BILLS = 40  # each bill sends an email; keeps one request under the server timeout
MONTH_RE = re.compile(r"^\d{4}-(0[1-9]|1[0-2])$")
MONTH_FORMATS = ("%Y-%m", "%b-%Y", "%b %Y", "%B %Y", "%B-%Y", "%m/%Y")

KINDS = ("rent", "electricity")

ALIASES = {
    "tenant": {"tenant", "tenantname", "username", "user", "room", "roomno", "roomnumber", "name"},
    "electricity": {"electricity", "electricitybill", "electricitybillamount", "electricityamount", "eb", "ebbill"},
    "month": {"month", "monthyear", "billmonth", "period"},
    "rent": {"rent", "rentamount"},
    "water": {"water", "waterbill", "wateramount"},
    "miscellaneous": {"miscellaneous", "misc", "miscellaneousamount", "maintenance", "other", "others"},
}

TEMPLATES = {
    "rent": ["Username", "Rent", "Water", "Miscellaneous"],
    "electricity": ["Username", "Electricity Bill Amount"],
}


class BulkError(ValueError):
    """The whole file is unusable (as opposed to a single bad row)."""


def check_kind(kind: Optional[str]) -> str:
    kind = (kind or "").lower()
    if kind not in KINDS:
        raise BulkError("Choose Rent or Electricity.")
    return kind


def _norm(header: Any) -> str:
    return re.sub(r"[^a-z0-9]", "", str(header or "").lower())


def _map_headers(headers: List[Any]) -> Dict[int, str]:
    mapped: Dict[int, str] = {}
    for idx, h in enumerate(headers):
        n = _norm(h)
        for field, names in ALIASES.items():
            if n in names and field not in mapped.values():
                mapped[idx] = field
                break
    return mapped


def _is_blank(v: Any) -> bool:
    return v is None or (isinstance(v, str) and not v.strip())


def read_rows(filename: str, data: bytes, kind: str) -> List[Tuple[int, Dict[str, Any]]]:
    """Returns [(spreadsheet row number, {field: raw value})] for non-empty rows."""
    kind = check_kind(kind)
    name = (filename or "").lower()
    if name.endswith(".xlsx"):
        try:
            from openpyxl import load_workbook
            sheet = load_workbook(io.BytesIO(data), read_only=True, data_only=True).worksheets[0]
            table = [list(r) for r in sheet.iter_rows(values_only=True)]
        except Exception:
            raise BulkError("Could not read this Excel file. Use the template and save it as .xlsx.")
    elif name.endswith(".csv"):
        try:
            text = data.decode("utf-8-sig")
        except UnicodeDecodeError:
            text = data.decode("latin-1")
        table = list(csv.reader(io.StringIO(text)))
    else:
        raise BulkError("Upload an Excel (.xlsx) or CSV file.")

    if not table:
        raise BulkError("The file is empty.")
    fields = set((mapped := _map_headers(table[0])).values())
    if "tenant" not in fields:
        raise BulkError("Missing the Username column (the first row must be the headings).")
    if kind == "rent":
        if "rent" not in fields:
            raise BulkError("The Rent upload needs a Rent column.")
        if "electricity" in fields:
            raise BulkError("This is the Rent upload, but the file has an Electricity column. Remove it, or use the Electricity upload.")
    else:
        if "electricity" not in fields:
            raise BulkError("The Electricity upload needs an Electricity Bill Amount column.")
        if fields & {"rent", "water", "miscellaneous"}:
            raise BulkError("This is the Electricity upload, but the file has Rent, Water or Miscellaneous columns. Remove them, or use the Rent upload.")

    rows = []
    for n, raw in enumerate(table[1:], start=2):
        record = {field: (raw[i] if i < len(raw) else None) for i, field in mapped.items()}
        if all(_is_blank(v) for v in record.values()):
            continue
        rows.append((n, record))
    if not rows:
        raise BulkError("The file has headings but no data rows.")
    return rows


def parse_amount(raw: Any) -> Optional[float]:
    if _is_blank(raw):
        return None
    if isinstance(raw, str):
        raw = re.sub(r"(?i)rs\.?|inr|₹|,|\s", "", raw)
    try:
        value = float(raw)
    except (TypeError, ValueError):
        raise ValueError(f"'{raw}' is not a valid amount")
    if math.isnan(value) or math.isinf(value) or value < 0:
        raise ValueError(f"'{raw}' is not a valid amount")
    return value


def parse_month(raw: Any) -> str:
    if isinstance(raw, (datetime, date)):
        return raw.strftime("%Y-%m")
    text = str(raw or "").strip()
    if MONTH_RE.match(text):
        return text
    for fmt in MONTH_FORMATS:
        try:
            return datetime.strptime(text, fmt).strftime("%Y-%m")
        except ValueError:
            continue
    raise ValueError(f"'{text}' is not a valid month (use YYYY-MM, e.g. 2026-09)")


def _entry(row, tenant, month, bill_type, amounts, status, message, bill=None) -> Dict[str, Any]:
    """`amounts` is {"rent","water","miscellaneous"} or {"electricity"} -- shown as
    separate columns in the preview (the total is worked out on the bill itself)."""
    return {"row": row, "tenant": tenant, "month": month, "billType": bill_type,
            **(amounts or {}), "status": status, "message": message, "_bill": bill}


def plan_rows(rows, kind: str, default_month: Optional[str], user_repo, bill_repo) -> List[Dict[str, Any]]:
    """Validates every row. Saves nothing. Each row yields exactly one bill of `kind`."""
    kind = check_kind(kind)
    bill_type = "RENT" if kind == "rent" else "ELECTRICITY"
    fallback_month = None
    if default_month:
        try:
            fallback_month = parse_month(default_month)
        except ValueError as exc:
            raise BulkError(str(exc))

    planned: List[Dict[str, Any]] = []
    seen = set()
    for row_no, rec in rows:
        tenant = str(rec.get("tenant") or "").strip()
        try:
            if not tenant:
                raise ValueError("Username is empty")
            user = user_repo.find_by_username(tenant)
            if user is None:
                raise ValueError(f"Tenant '{tenant}' not found")
            if user.role != "TENANT":
                raise ValueError(f"'{tenant}' is not a tenant account")

            if not _is_blank(rec.get("month")):
                month = parse_month(rec.get("month"))
            elif fallback_month:
                month = fallback_month
            else:
                raise ValueError("No month: add a Month column or choose the month above")

            if kind == "rent":
                rent = parse_amount(rec.get("rent"))
                if rent is None:
                    raise ValueError("Rent is empty")
                water = parse_amount(rec.get("water")) or 0.0
                misc = parse_amount(rec.get("miscellaneous")) or 0.0
                amount = {"rent": rent, "water": water, "miscellaneous": misc}
                bill = TenantBill(tenant_name=tenant, month_year=month, bill_type="RENT",
                                  rent=rent, water=water, electricity=None, miscellaneous=misc, paid=False)
            else:
                elec = parse_amount(rec.get("electricity"))
                if elec is None or elec <= 0:
                    raise ValueError("Electricity amount is empty or zero")
                amount = {"electricity": elec}
                bill = TenantBill(tenant_name=tenant, month_year=month, bill_type="ELECTRICITY",
                                  rent=None, water=None, electricity=elec, miscellaneous=None, paid=False)
        except ValueError as exc:
            planned.append(_entry(row_no, tenant, None, bill_type, None, "error", str(exc)))
            continue

        key = (tenant, month)
        if key in seen:
            planned.append(_entry(row_no, tenant, month, bill_type, amount, "duplicate", "Repeated in this file - skipped"))
        elif bill_repo.find_by_tenant_name_and_month(tenant, month, bill_type) is not None:
            planned.append(_entry(row_no, tenant, month, bill_type, amount, "duplicate", "Bill already exists - skipped"))
        else:
            planned.append(_entry(row_no, tenant, month, bill_type, amount, "ok", "Ready to import", bill))
        seen.add(key)

    ok_count = sum(1 for p in planned if p["status"] == "ok")
    if ok_count > MAX_BILLS:
        raise BulkError(f"{ok_count} bills in one file is too many. Upload at most {MAX_BILLS} at a time.")
    return planned


def import_planned(planned: List[Dict[str, Any]], bill_service) -> None:
    """Saves the valid bills (one email each, via add_bill). Updates status in place."""
    for p in planned:
        if p["status"] != "ok":
            continue
        try:
            bill_service.add_bill(p["_bill"])
            p["status"], p["message"] = "added", "Added, email sent"
        except ValueError as exc:
            p["status"], p["message"] = "duplicate", str(exc)
        except Exception:
            p["status"], p["message"] = "error", "Could not save this bill"


def public_rows(planned: List[Dict[str, Any]]) -> List[Dict[str, Any]]:
    return [{k: v for k, v in p.items() if k != "_bill"} for p in planned]


def summarize(planned: List[Dict[str, Any]]) -> Dict[str, int]:
    counts: Dict[str, int] = {}
    for p in planned:
        counts[p["status"]] = counts.get(p["status"], 0) + 1
    return counts


def build_template(kind: str) -> bytes:
    from openpyxl import Workbook
    headers = TEMPLATES[check_kind(kind)]
    wb = Workbook()
    ws = wb.active
    ws.title = "Bills"
    ws.append(headers)
    for col in range(1, len(headers) + 1):
        ws.column_dimensions[chr(64 + col)].width = 26
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()
